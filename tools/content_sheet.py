#!/usr/bin/env python3
"""Planilha mestre de frases <-> JSONs do app.

Uso (rodar na raiz do projeto):
  python3 tools/content_sheet.py export        # JSONs -> planilha (só na criação; com planilha existente exige --force)
  python3 tools/content_sheet.py check         # valida a planilha, não grava nada
  python3 tools/content_sheet.py build         # valida e gera os JSONs (só grava se não houver erro)
  python3 tools/content_sheet.py preview       # mostra o que mudou desde a última publicação
  python3 tools/content_sheet.py publish       # build + gera pacote remoto em dist/content/ (versão +1)
  python3 tools/content_sheet.py publish --disable   # manifesto com enabled=false (app volta às frases embutidas)

Regra: a planilha é a fonte da verdade. Os JSONs em Content/messages/ são gerados; não edite à mão.
Campos que não são frases (perguntas do diário, inline_contexto, limites, revisão) continuam
nos JSONs e são preservados pelo build.
"""
import json, re, sys
from collections import OrderedDict
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("Falta o openpyxl: pip3 install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
MSG_DIR = ROOT / "Content" / "messages"
XLSX = ROOT / "Content" / "Frases_Fudoshin.xlsx"
NUMBERS = ROOT / "Content" / "Frases_Fudoshin.numbers"   # se existir, é a fonte da verdade

STATUS = ["ativa", "rascunho", "removida"]
META_COLS = ["id", "autor", "obra", "ref", "tipo", "gatilhos",
             "base_dominio_publico", "texto_original", "verificacao"]
TEXT_FIELDS = [("texto", "texto"), ("curto", "texto_curto"),
               ("inline", "texto_inline"), ("reflexao", "reflexao")]
TIPOS = ["citacao", "parafrase", "propria"]
ID_RE = re.compile(r"^[a-z0-9]+(-[a-z0-9]+)*$")

FONT = "Arial"
HEAD_FILL = PatternFill("solid", start_color="1F2A44")
META_FILL = PatternFill("solid", start_color="EEF1F6")
RED_FILL = PatternFill("solid", start_color="F8C9C4")


def load(p):
    return json.loads(p.read_text(encoding="utf-8"), object_pairs_hook=OrderedDict)


def dump(p, data):
    p.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def langs_of(base):
    return list(base["idiomas_lancamento"])


def text_cols(langs):
    return [f"{short}_{l}" for l in langs for short, _ in TEXT_FIELDS]


# ---------------------------------------------------------------- export
def export():
    if (XLSX.exists() or NUMBERS.exists()) and "--force" not in sys.argv:
        sys.exit("A planilha já existe. 'export' recria a partir dos JSONs e APAGA rascunhos, removidas e "
                 "edições não convertidas. Se é isso mesmo: python3 tools/content_sheet.py export --force")
    base = load(MSG_DIR / "base.json")
    langs = langs_of(base)
    lang = {l: load(MSG_DIR / f"{l}.json") for l in langs}
    lim = base["limites_widget"]["padrao"]

    wb = openpyxl.Workbook()
    guide = wb.active
    guide.title = "Como usar"
    ws = wb.create_sheet("Mensagens")
    wa = wb.create_sheet("Autores")
    wo = wb.create_sheet("Obras")
    wg = wb.create_sheet("Gatilhos")

    # Mensagens
    cols = ["status"] + META_COLS + text_cols(langs)
    ws.append(cols)
    for m in base["mensagens"]:
        row = ["ativa"]
        for c in META_COLS:
            v = m.get(c, "")
            row.append(", ".join(v) if c == "gatilhos" else v)
        for l in langs:
            t = lang[l]["mensagens"].get(m["id"], {})
            extra = set(t) - {k for _, k in TEXT_FIELDS}
            if extra:
                sys.exit(f"{l}.json/{m['id']}: campos desconhecidos {extra}; ajuste o script antes de exportar")
            row += [t.get(k, "") for _, k in TEXT_FIELDS]
        ws.append(row)
    style_table(ws, len(cols))
    widths = {"status": 11, "id": 18, "autor": 16, "obra": 18, "ref": 8, "tipo": 11, "gatilhos": 26,
              "base_dominio_publico": 26, "texto_original": 30, "verificacao": 30}
    for i, c in enumerate(cols, 1):
        w = widths.get(c) or (16 if c.startswith("inline") else 30 if c.startswith("curto") else 45)
        ws.column_dimensions[get_column_letter(i)].width = w
        if i <= len(META_COLS) + 1:
            for r in range(2, ws.max_row + 1):
                ws.cell(r, i).fill = META_FILL
    ws.freeze_panes = "C2"
    last = 1000
    for i, c in enumerate(cols, 1):
        L = get_column_letter(i)
        rng = f"{L}2:{L}{last}"
        if c.startswith("curto_"):
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f"LEN({L}2)>{lim['texto_curto_max']}"], fill=RED_FILL))
        if c.startswith("inline_"):
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f"LEN({L}2)>{lim['texto_inline_max']}"], fill=RED_FILL))
        if c.startswith(("texto_", "curto_", "inline_", "reflexao_")) and c != "texto_original" or c in ("autor", "obra", "tipo", "gatilhos"):
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND($A2="ativa",LEN(TRIM({L}2))=0)'], fill=RED_FILL))
    dv = lambda f: DataValidation(type="list", formula1=f, allow_blank=False, showErrorMessage=True)
    col = {c: get_column_letter(i) for i, c in enumerate(cols, 1)}
    for name, f in (("status", '"ativa,rascunho,removida"'), ("tipo", '"citacao,parafrase,propria"'),
                    ("autor", f"=Autores!$A$2:$A${last}"),
                    ("obra", f"=Obras!$A$2:$A${last}")):
        d = dv(f); ws.add_data_validation(d); d.add(f"{col[name]}2:{col[name]}{last}")

    # Autores / Obras
    for sheet, key in ((wa, "autores"), (wo, "obras")):
        sheet.append(["chave"] + langs)
        for k, names in base[key].items():
            sheet.append([k] + [names.get(l, "") for l in langs])
        style_table(sheet, 1 + len(langs))
        sheet.column_dimensions["A"].width = 24
        for i in range(2, 2 + len(langs)):
            sheet.column_dimensions[get_column_letter(i)].width = 34
        sheet.freeze_panes = "B2"

    # Gatilhos (referência; a lista válida vem de base.json)
    desc = {"marco_prova": "30, 7, 1 e 0 dias antes da prova", "pos_treino": "até 3h depois de um treino",
            "sono_ruim": "dormiu menos que a meta - 1h (manhã)", "vespera_chave": "véspera de longão ou treino forte",
            "descanso": "dia sem treino planejado", "retorno": "volta após 5+ dias parado", "geral": "qualquer dia"}
    wg.append(["gatilho", "quando aparece"])
    for g in base["gatilhos"]:
        wg.append([g, desc.get(g, "")])
    style_table(wg, 2)
    wg.column_dimensions["A"].width = 18; wg.column_dimensions["B"].width = 44

    # Como usar
    lines = [
        ("Planilha mestre de frases do Fudōshin", True),
        ("Esta planilha é a fonte da verdade. Os arquivos JSON do app são gerados a partir dela.", False),
        ("", False),
        ("Status de cada linha", True),
        ("ativa = entra no app. Precisa estar completa nos 3 idiomas.", False),
        ("rascunho = não entra no app. Pode estar incompleta (ex.: só em português, aguardando tradução).", False),
        ("removida = não entra no app. Fica na planilha como histórico, para não ser adicionada de novo.", False),
        ("", False),
        ("Onde editar", True),
        ("Aba Mensagens: uma linha por frase. Colunas cinza = dados da frase; colunas brancas = textos por idioma.", False),
        ("Abas Autores e Obras: nome exibido em cada idioma. Cadastre aqui antes de usar na aba Mensagens.", False),
        ("Aba Gatilhos: só referência. Para criar um gatilho novo é preciso mudar o código do app.", False),
        ("", False),
        ("Regras de preenchimento", True),
        ("id: minúsculas, números e hífen, único (ex.: sen-ira-3). Nunca mude o id de uma frase publicada.", False),
        ("tipo propria: frase original nossa (sem ref). tipo: citacao (trecho de obra em domínio público, ref obrigatória) ou parafrase (ideia com nossas palavras; app mostra 'Inspirado em').", False),
        ("gatilhos: separados por vírgula, ex.: marco_prova, vespera_chave", False),
        (f"curto_*: até {lim['texto_curto_max']} caracteres (widget e notificação). inline_*: até {lim['texto_inline_max']}. Célula fica vermelha se passar.", False),
        ("Célula vermelha em coluna obrigatória = está vazia.", False),
        ("verificacao: deixe 'pendente — ...' até alguém conferir a fonte; depois escreva quem conferiu e a data.", False),
        ("Direitos: siga Content/DIREITOS.md (domínio público, traduções próprias, autores vivos só como paráfrase).", False),
        ("", False),
        ("Exemplo de linha", True),
        ("id: sen-brevidade-1 | autor: seneca | obra: cartas_lucilio | ref: 1.1 | tipo: citacao | gatilhos: geral, descanso", False),
        ("", False),
        ("Depois de editar (no Terminal, na pasta AtletaEstoico)", True),
        ("python3 tools/content_sheet.py check   → só valida e lista os erros", False),
        ("python3 tools/content_sheet.py build   → valida e gera os JSONs (nada é gravado se houver erro)", False),
        ("Depois: rodar os testes no Xcode e fazer commit da planilha e dos JSONs juntos.", False),
    ]
    for i, (t, bold) in enumerate(lines, 1):
        c = guide.cell(i, 1, t)
        c.font = Font(name=FONT, bold=bold, size=13 if bold else 11)
    guide.column_dimensions["A"].width = 120
    wb.active = 1
    wb.save(XLSX)
    print(f"Planilha gerada: {XLSX.relative_to(ROOT)} ({len(base['mensagens'])} mensagens)")


def style_table(ws, ncols):
    for r in ws.iter_rows():
        for c in r:
            c.font = Font(name=FONT, size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for i in range(1, ncols + 1):
        h = ws.cell(1, i)
        h.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        h.fill = HEAD_FILL
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{max(ws.max_row, 2)}"


# ---------------------------------------------------------------- read + validate
def cell(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def open_source():
    """Numbers tem prioridade (é onde o JC edita). Sem ele, usa o .xlsx."""
    if NUMBERS.exists():
        try:
            from numbers_parser import Document
        except ImportError:
            sys.exit("Para ler o arquivo .numbers instale: pip3 install --user numbers-parser")
        doc = Document(str(NUMBERS))
        sheets = {s.name: s.tables[0].rows(values_only=True) for s in doc.sheets if s.tables}
        return NUMBERS.name, sheets
    if not XLSX.exists():
        sys.exit("Nenhuma planilha encontrada em Content/. Rode 'export' primeiro.")
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    return XLSX.name, {n: list(wb[n].iter_rows(values_only=True)) for n in wb.sheetnames}


def read_sheet(wb, name):
    if name not in wb:
        raise SystemExit(f"Aba '{name}' não encontrada na planilha")
    rows = list(wb[name])
    head = [cell(h) for h in rows[0]]
    out = []
    for n, r in enumerate(rows[1:], 2):
        d = {h: cell(v) for h, v in zip(head, r) if h}
        if any(d.values()):
            out.append((n, d))
    return head, out


def collect(base):
    langs = langs_of(base)
    lim = base["limites_widget"]["padrao"]
    src, wb = open_source()
    print(f"Lendo {src}")
    errors, warns = [], []

    def names(sheet):
        head, rows = read_sheet(wb, sheet)
        miss = [l for l in langs if l not in head]
        if miss:
            errors.append(f"{sheet}: faltam colunas {miss}")
        res = OrderedDict()
        for n, d in rows:
            k = d.get("chave", "")
            if not k:
                errors.append(f"{sheet} linha {n}: chave vazia"); continue
            if k in res:
                errors.append(f"{sheet} linha {n}: chave duplicada '{k}'")
            for l in langs:
                if not d.get(l):
                    errors.append(f"{sheet} linha {n} ({k}): nome em {l} vazio")
            res[k] = OrderedDict((l, d.get(l, "")) for l in langs)
        return res

    autores, obras = names("Autores"), names("Obras")
    head, rows = read_sheet(wb, "Mensagens")
    need = ["status"] + META_COLS + text_cols(langs)
    miss = [c for c in need if c not in head]
    if miss:
        errors.append(f"Mensagens: faltam colunas {miss}")
        return None, None, None, errors, warns

    valid_trig = set(base["gatilhos"])
    msgs, texts, seen = [], {l: OrderedDict() for l in langs}, {}
    pend = 0
    skipped = {"rascunho": 0, "removida": 0}
    for n, d in rows:
        mid = d["id"]
        where = f"Mensagens linha {n} ({mid or 'sem id'})"
        st = d["status"].lower()
        if st not in STATUS:
            errors.append(f"{where}: status deve ser ativa, rascunho ou removida"); continue
        if mid and mid in seen:
            errors.append(f"{where}: id repetido (também na linha {seen[mid]})")
        if mid:
            seen.setdefault(mid, n)
        if st != "ativa":
            skipped[st] += 1
            if st == "rascunho":
                faltam = [f"{s}_{l}" for l in langs for s, _ in TEXT_FIELDS if not d[f"{s}_{l}"]]
                if faltam:
                    warns.append(f"{where}: rascunho, falta {', '.join(faltam)}")
            continue
        if not ID_RE.match(mid):
            errors.append(f"{where}: id inválido (use minúsculas, números e hífen)")
        if d["autor"] not in autores:
            errors.append(f"{where}: autor '{d['autor']}' não está na aba Autores")
        if d["obra"] not in obras:
            errors.append(f"{where}: obra '{d['obra']}' não está na aba Obras")
        if d["tipo"] not in TIPOS:
            errors.append(f"{where}: tipo deve ser citacao, parafrase ou propria")
        if d["tipo"] == "citacao" and not d["ref"]:
            errors.append(f"{where}: citação sem ref (livro.capítulo) — DIREITOS.md regra 3")
        if d["tipo"] == "citacao" and not d["base_dominio_publico"]:
            errors.append(f"{where}: citação sem base_dominio_publico")
        trig = [t.strip() for t in d["gatilhos"].split(",") if t.strip()]
        if not trig:
            errors.append(f"{where}: sem gatilhos")
        for t in trig:
            if t not in valid_trig:
                errors.append(f"{where}: gatilho '{t}' não existe (válidos: {', '.join(base['gatilhos'])})")
        if len(set(trig)) != len(trig):
            errors.append(f"{where}: gatilho repetido")
        if d["verificacao"].lower().startswith("pendente"):
            pend += 1
        m = OrderedDict()
        for c in META_COLS:
            if c == "gatilhos":
                m[c] = trig
            elif c in ("id", "autor", "obra", "ref", "tipo"):
                m[c] = d[c]          # sempre presentes (ref pode ser "")
            elif d[c]:
                m[c] = d[c]          # opcionais: só se preenchidos
        msgs.append(m)
        for l in langs:
            t = OrderedDict()
            for short, key in TEXT_FIELDS:
                v = d[f"{short}_{l}"]
                if not v:
                    errors.append(f"{where}: {short}_{l} vazio")
                t[key] = v
            if len(t["texto_curto"]) > lim["texto_curto_max"]:
                errors.append(f"{where}: curto_{l} com {len(t['texto_curto'])} caracteres (máx. {lim['texto_curto_max']})")
            if len(t["texto_inline"]) > lim["texto_inline_max"]:
                errors.append(f"{where}: inline_{l} com {len(t['texto_inline'])} caracteres (máx. {lim['texto_inline_max']})")
            if "  " in "".join(t.values()):
                warns.append(f"{where} [{l}]: espaço duplo no texto")
            texts[l][mid] = t
    if not msgs:
        errors.append("Mensagens: nenhuma frase encontrada")
    used_a = {m["autor"] for m in msgs}; used_o = {m["obra"] for m in msgs}
    for k in autores:
        if k not in used_a: warns.append(f"Autor '{k}' cadastrado sem nenhuma frase")
    for k in obras:
        if k not in used_o: warns.append(f"Obra '{k}' cadastrada sem nenhuma frase")
    for t in valid_trig:
        cnt = sum(t in m["gatilhos"] for m in msgs)
        if cnt < 3: warns.append(f"Gatilho '{t}' tem só {cnt} frase(s): risco de repetição")
    if any(skipped.values()):
        warns.append(f"Fora do app: {skipped['rascunho']} rascunho(s), {skipped['removida']} removida(s)")
    if pend:
        warns.append(f"{pend} frase(s) com verificação pendente — conferir antes de publicar")
    return (autores, obras, msgs), texts, langs, errors, warns


def run(write):
    base = load(MSG_DIR / "base.json")
    meta, texts, langs, errors, warns = collect(base)
    for w in warns:
        print("AVISO:", w)
    if errors:
        print(f"\n{len(errors)} ERRO(S) — nada foi gravado:")
        for e in errors:
            print("  -", e)
        sys.exit(1)
    autores, obras, msgs = meta
    print(f"OK: {len(msgs)} mensagens, {len(autores)} autores, {len(obras)} obras, idiomas {', '.join(langs)}")
    if not write:
        return
    base["autores"], base["obras"], base["mensagens"] = autores, obras, msgs
    dump(MSG_DIR / "base.json", base)
    for l in langs:
        p = MSG_DIR / f"{l}.json"
        data = load(p)
        data["mensagens"] = texts[l]
        dump(p, data)
    print("JSONs gravados em Content/messages/. Rode os testes no Xcode antes do commit.")


DIST = ROOT / "dist" / "content"
LIMIT_BYTES = 2_000_000


def preview():
    """Compara a planilha com o último pacote publicado (sem gravar nada)."""
    base = load(MSG_DIR / "base.json")
    meta, texts, langs, errors, warns = collect(base)
    if errors:
        print(f"{len(errors)} ERRO(S) na planilha — corrija antes de publicar:")
        for e in errors: print("  -", e)
        sys.exit(1)
    autores, obras, msgs = meta
    new_ids = {m["id"]: m for m in msgs}
    mf = DIST / "manifest.json"
    if not mf.exists():
        print(f"Primeira publicação: {len(new_ids)} mensagens."); return
    old = load(DIST / load(mf)["url"])
    old_ids = {m["id"]: m for m in old["base"]["mensagens"]}
    pt = langs[0]
    added = [i for i in new_ids if i not in old_ids]
    removed = [i for i in old_ids if i not in new_ids]
    changed = [i for i in new_ids if i in old_ids and (
        new_ids[i] != old_ids[i] or any(texts[l][i] != old["locales"][l]["mensagens"].get(i) for l in langs))]
    print(f"Último publicado: v{old['version']} ({len(old_ids)} mensagens) → agora: {len(new_ids)} mensagens")
    for title, ids in (("NOVAS", added), ("ALTERADAS", changed)):
        print(f"\n{title} ({len(ids)}):")
        for i in ids: print(f"  + {i}: {texts[pt][i]['texto_curto']}")
    print(f"\nREMOVIDAS ({len(removed)}):")
    for i in removed: print(f"  - {i}")
    if not (added or removed or changed):
        print("\nNada mudou desde a última publicação."); sys.exit(3)


def publish(disable=False):
    import hashlib, datetime
    run(write=True)
    base = load(MSG_DIR / "base.json")
    langs = langs_of(base)
    DIST.mkdir(parents=True, exist_ok=True)
    last = 0
    mf = DIST / "manifest.json"
    if mf.exists():
        last = load(mf).get("version", 0)
    version = max(last, base.get("versao_conteudo", 0)) + 1
    base["versao_conteudo"] = version
    dump(MSG_DIR / "base.json", base)          # o próximo build do app embute a mesma versão
    pkg = OrderedDict(schema=base["versao"], version=version,
                      generatedAt=datetime.datetime.now(datetime.timezone.utc).isoformat(timespec="seconds"),
                      base=base,
                      locales=OrderedDict((l, load(MSG_DIR / f"{l}.json")) for l in langs))
    raw = json.dumps(pkg, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    if len(raw) > LIMIT_BYTES:
        sys.exit(f"Pacote com {len(raw)} bytes passa do limite de {LIMIT_BYTES}")
    name = f"content-v{version}.json"
    (DIST / name).write_bytes(raw)
    manifest = OrderedDict(schema=base["versao"], version=version, url=name,
                           sha256=hashlib.sha256(raw).hexdigest(), bytes=len(raw),
                           enabled=not disable, minAppBuild=1)
    dump(mf, manifest)
    n = len(base["mensagens"])
    print(f"Pacote v{version} gerado em dist/content/ ({n} mensagens, {len(raw)//1024} KB)"
          + ("  [DESATIVADO]" if disable else ""))
    print("Próximo passo: copie dist/content/* para o repositório do GitHub Pages e faça push.")


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else ""
    if cmd == "export": export()
    elif cmd == "check": run(write=False)
    elif cmd == "build": run(write=True)
    elif cmd == "publish": publish(disable="--disable" in sys.argv)
    elif cmd == "preview": preview()
    else: sys.exit(__doc__)
